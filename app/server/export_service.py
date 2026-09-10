import os
import re
import traceback
import unicodedata
from xml.sax.saxutils import escape as xml_escape

import config
from file_extractors import extract_text_from_file
from logging_utils import log_message

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
except ImportError:
    colors = None
    TA_LEFT = None
    A4 = None
    ParagraphStyle = None
    getSampleStyleSheet = None
    mm = None
    Paragraph = None
    SimpleDocTemplate = None
    Spacer = None


SECTION_HEADINGS = ("summary", "goals", "skills", "steps", "materials", "assessment", "reflection")


def normalize_filename_piece(value, default="lesson-plan", max_length=80):
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-").lower()
    text = text or default
    if max_length and len(text) > max_length:
        text = text[:max_length].rstrip("-") or default
    return text


def normalize_text_list(value):
    if not isinstance(value, list):
        return []
    cleaned = []
    for item in value:
        text = str(item or "").strip()
        if text:
            cleaned.append(text)
    return cleaned


def normalize_source_document(source_document):
    if isinstance(source_document, str):
        source_document = {
            "path": source_document,
        }
    if not isinstance(source_document, dict):
        return None

    normalized = dict(source_document)
    path = str(normalized.get("path", "") or normalized.get("storedPath", "") or normalized.get("filePath", "")).strip()
    if path:
        normalized["path"] = path
    filename = str(normalized.get("filename", "") or normalized.get("originalFilename", "") or os.path.basename(path)).strip()
    if filename:
        normalized["filename"] = filename
    stored_name = str(normalized.get("storedFilename", "")).strip()
    if stored_name:
        normalized["storedFilename"] = stored_name
    return normalized if path or filename else None


def find_uploaded_file_path(filename):
    filename = os.path.basename(str(filename or "").strip())
    if not filename:
        return ""

    direct_path = os.path.join(config.LESSONPLANS_ROOT, filename)
    if os.path.exists(direct_path):
        return direct_path

    matches = []
    for root, _, files in os.walk(config.LESSONPLANS_ROOT):
        if filename in files:
            candidate = os.path.join(root, filename)
            try:
                mtime = os.path.getmtime(candidate)
            except OSError:
                mtime = 0
            matches.append((mtime, candidate))

    if matches:
        matches.sort(key=lambda item: item[0], reverse=True)
        return matches[0][1]

    return ""


def resolve_source_document(context):
    context = context if isinstance(context, dict) else {}
    payload = context.get("payload") if isinstance(context.get("payload"), dict) else {}
    state = context.get("state") if isinstance(context.get("state"), dict) else {}

    candidates = [
        payload.get("sourceDocument"),
        payload.get("source_document"),
        payload.get("lessonPlanPath"),
        payload.get("lesson_plan_path"),
        state.get("sourceDocument"),
    ]
    if isinstance(state.get("payload"), dict):
        candidates.extend([
            state["payload"].get("sourceDocument"),
            state["payload"].get("source_document"),
            state["payload"].get("lessonPlanPath"),
            state["payload"].get("lesson_plan_path"),
        ])

    for candidate in candidates:
        normalized = normalize_source_document(candidate)
        if not normalized:
            continue
        path = normalized.get("path", "")
        if path and os.path.exists(path):
            return normalized

    filename_candidates = [
        payload.get("filename"),
        payload.get("sourceFilename"),
        payload.get("source_filename"),
        payload.get("lessonPlanFilename"),
        payload.get("lesson_plan_filename"),
    ]
    meta = payload.get("meta") if isinstance(payload.get("meta"), dict) else {}
    filename_candidates.extend([
        meta.get("filename"),
        meta.get("sourceFilename"),
        meta.get("source_filename"),
    ])
    if isinstance(state.get("payload"), dict):
        state_payload = state["payload"]
        filename_candidates.extend([
            state_payload.get("filename"),
            state_payload.get("sourceFilename"),
            state_payload.get("source_filename"),
        ])

    for filename in filename_candidates:
        found_path = find_uploaded_file_path(filename)
        if found_path:
            return {
                "path": found_path,
                "filename": os.path.basename(found_path),
            }

    return None


def build_export_context(payload, state_loader=None):
    payload = payload if isinstance(payload, dict) else {}
    session_id = str(payload.get("sessionId", "")).strip()
    state = state_loader(session_id) if session_id and state_loader else None
    state = state if isinstance(state, dict) else {}
    state_payload = state.get("payload") if isinstance(state.get("payload"), dict) else {}

    draft = None
    if isinstance(state.get("draft"), dict):
        draft = state.get("draft")
    if not draft and isinstance(payload.get("draft"), dict):
        draft = payload.get("draft")
    if not draft and isinstance(state_payload.get("draft"), dict):
        draft = state_payload.get("draft")
    if not isinstance(draft, dict):
        draft = {}

    conversation = None
    if isinstance(state.get("conversation"), list):
        conversation = state.get("conversation")
    if conversation is None and isinstance(payload.get("conversation"), list):
        conversation = payload.get("conversation")
    if conversation is None and isinstance(state_payload.get("conversation"), list):
        conversation = state_payload.get("conversation")
    if not isinstance(conversation, list):
        conversation = []

    source_document = resolve_source_document({
        "payload": payload,
        "state": state,
    })

    if not source_document and isinstance(state_payload, dict):
        source_document = resolve_source_document({
            "payload": state_payload,
            "state": state,
        })

    source_text = ""
    if source_document and source_document.get("path") and os.path.exists(source_document["path"]):
        source_text = extract_text_from_file(source_document["path"])

    if not source_text:
        source_text = str(draft.get("summary") or "").strip()

    if not source_document and source_text:
        source_document = {
            "filename": "lesson-plan",
            "path": "",
        }

    return {
        "sessionId": session_id,
        "state": state,
        "payload": payload,
        "draft": draft,
        "conversation": conversation,
        "sourceDocument": source_document,
        "sourceText": source_text,
    }


def collect_export_changes(context):
    context = context if isinstance(context, dict) else {}
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else {}
    state = context.get("state") if isinstance(context.get("state"), dict) else {}
    payload = context.get("payload") if isinstance(context.get("payload"), dict) else {}
    result = state.get("result") if isinstance(state.get("result"), dict) else {}

    raw_changes = []
    if isinstance(draft.get("changes"), list) and draft.get("changes"):
        raw_changes = draft.get("changes")
    elif isinstance(result.get("changes_made"), list) and result.get("changes_made"):
        raw_changes = result.get("changes_made")
    elif isinstance(payload.get("changes_made"), list) and payload.get("changes_made"):
        raw_changes = payload.get("changes_made")

    changes = []
    for item in raw_changes:
        item = item if isinstance(item, dict) else {}
        title = str(item.get("title") or item.get("section") or item.get("kind") or "Change").strip()
        reason = str(item.get("reason") or item.get("summary") or item.get("note") or "").strip()
        location = str(item.get("location") or item.get("section") or item.get("source") or "").strip()
        before_value = str(item.get("beforeValue") or item.get("before") or "").strip()
        after_value = str(item.get("afterValue") or item.get("after") or "").strip()
        source = str(item.get("source") or item.get("section") or "").strip()
        changes.append({
            "title": title,
            "reason": reason,
            "location": location,
            "beforeValue": before_value,
            "afterValue": after_value,
            "source": source,
        })

    return changes


def get_final_plan_context(context):
    context = context if isinstance(context, dict) else {}
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else {}
    source_document = context.get("sourceDocument") if isinstance(context.get("sourceDocument"), dict) else {}
    source_text = str(context.get("sourceText") or "").strip()
    title = str(draft.get("title") or source_document.get("filename") or "Lesson plan").strip()
    summary = str(draft.get("summary") or "").strip()
    meta = draft.get("meta") if isinstance(draft.get("meta"), dict) else {}
    goals = normalize_text_list(draft.get("goals"))
    skills = normalize_text_list(draft.get("skills"))
    materials = normalize_text_list(draft.get("materials"))
    assessment = normalize_text_list(draft.get("assessment"))
    steps = draft.get("steps") if isinstance(draft.get("steps"), list) else []
    reflection = str(draft.get("reflection") or "").strip()
    changes = collect_export_changes(context)
    conversation = context.get("conversation") if isinstance(context.get("conversation"), list) else []

    return {
        "title": title,
        "summary": summary,
        "meta": meta,
        "goals": goals,
        "skills": skills,
        "steps": steps,
        "materials": materials,
        "assessment": assessment,
        "reflection": reflection,
        "changes": changes,
        "conversation": conversation,
        "sourceDocument": source_document,
        "sourceText": source_text,
    }


def build_export_metadata_lines(context):
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else {}
    meta = draft.get("meta") if isinstance(draft.get("meta"), dict) else {}
    source_document = context.get("sourceDocument") if isinstance(context.get("sourceDocument"), dict) else {}
    lines = []
    for label, value in [
        ("Country", meta.get("country", "")),
        ("Subject(s)", ", ".join(meta.get("subjects", [])) if isinstance(meta.get("subjects"), list) else meta.get("subjects", "")),
        ("Grade range", meta.get("gradeRange", "")),
        ("Age range", meta.get("ageRange", "")),
        ("Focus", meta.get("focus", "")),
        ("Original file", source_document.get("filename", "")),
    ]:
        value_text = str(value or "").strip()
        if value_text:
            lines.append(f"{label}: {value_text}")
    return lines


def normalize_heading_key(value):
    text = str(value or "").strip().lower()
    text = re.sub(r"^\s*\d+\s*[\.\)]\s*", "", text)
    text = text.replace("subject(s)", "subjects")
    text = text.replace("/", " ")
    text = re.sub(r"[^a-z0-9]+", " ", text).strip()
    return text


def split_text_to_paragraphs(text):
    blocks = []
    for raw_block in str(text or "").splitlines():
        line = raw_block.rstrip()
        if not line:
            blocks.append("")
            continue
        blocks.append(line)
    return blocks


def replace_first_non_empty_line(lines, new_value):
    if not lines:
        return [new_value] if new_value else []
    updated = list(lines)
    for index, line in enumerate(updated):
        if str(line or "").strip():
            if new_value:
                updated[index] = new_value
            return updated
    if new_value:
        updated.insert(0, new_value)
    return updated


def replace_metadata_lines(lines, metadata_lines):
    updated = list(lines)
    for meta_line in metadata_lines:
        label = meta_line.split(":", 1)[0].strip().lower() if ":" in meta_line else ""
        if not label:
            continue
        replaced = False
        for index, line in enumerate(updated):
            normalized = str(line or "").strip().lower()
            if normalized.startswith(label + ":") or normalized == label:
                updated[index] = meta_line
                replaced = True
                break
        if not replaced:
            insert_at = 1 if updated else 0
            while insert_at < len(updated) and not str(updated[insert_at]).strip():
                insert_at += 1
            updated.insert(insert_at, meta_line)
            updated.insert(insert_at + 1, "")
    return updated


def apply_inline_replacements(lines, changes):
    updated = list(lines)
    changes = changes if isinstance(changes, list) else []
    for change in changes:
        before_value = str((change or {}).get("beforeValue") or "").strip()
        after_value = str((change or {}).get("afterValue") or "").strip()
        if not before_value or not after_value:
            continue
        updated = [line.replace(before_value, after_value) for line in updated]
    return updated


def replace_section_block(lines, heading, content_lines):
    content_lines = list(content_lines or [])
    normalized_heading = normalize_heading_key(heading)
    updated = list(lines)
    heading_index = None

    for index, line in enumerate(updated):
        if normalize_heading_key(line) == normalized_heading:
            heading_index = index
            break

    if heading_index is None:
        if content_lines:
            if updated and str(updated[-1]).strip():
                updated.append("")
            updated.append(heading)
            updated.extend(content_lines)
        return updated

    end_index = len(updated)
    for index in range(heading_index + 1, len(updated)):
        candidate = updated[index]
        candidate_key = normalize_heading_key(candidate)
        if candidate_key in SECTION_HEADINGS and candidate_key != normalized_heading:
            end_index = index
            break

    updated = updated[: heading_index + 1] + content_lines + updated[end_index:]
    return updated


def build_final_export_lines(context):
    final_plan = get_final_plan_context(context)
    source_text = str(final_plan.get("sourceText") or "").strip()
    lines = split_text_to_paragraphs(source_text) if source_text else []
    lines = apply_inline_replacements(lines, final_plan.get("changes"))
    lines = replace_first_non_empty_line(lines, final_plan["title"])
    lines = replace_metadata_lines(lines, build_export_metadata_lines(context))
    lines = replace_section_block(lines, "Summary", [final_plan["summary"]] if final_plan["summary"] else [])
    lines = replace_section_block(lines, "Goals", [f"- {item}" for item in normalize_text_list(final_plan.get("goals"))])
    lines = replace_section_block(lines, "Skills", [f"- {item}" for item in normalize_text_list(final_plan.get("skills"))])

    step_lines = []
    steps = final_plan.get("steps") if isinstance(final_plan.get("steps"), list) else []
    for index, step in enumerate(steps, start=1):
        step = step if isinstance(step, dict) else {}
        title = str(step.get("title") or f"Step {index}").strip()
        duration = str(step.get("duration") or "").strip()
        description = str(step.get("description") or "").strip()
        status = str(step.get("status") or "").strip()
        note = str(step.get("note") or "").strip()
        step_lines.append(f"{index}. {title}" + (f" ({duration})" if duration else ""))
        if description:
            step_lines.append(description)
        if status:
            step_lines.append(f"Status: {status}")
        if note:
            step_lines.append(f"Note: {note}")
        step_lines.append("")
    if step_lines and step_lines[-1] == "":
        step_lines.pop()
    lines = replace_section_block(lines, "Steps", step_lines)
    lines = replace_section_block(lines, "Materials", [f"- {item}" for item in normalize_text_list(final_plan.get("materials"))])
    lines = replace_section_block(lines, "Assessment", [f"- {item}" for item in normalize_text_list(final_plan.get("assessment"))])
    lines = replace_section_block(lines, "Reflection", [final_plan["reflection"]] if final_plan["reflection"] else [])

    lines = [line for line in lines if line is not None]
    while lines and not str(lines[0]).strip():
        lines.pop(0)
    return lines, final_plan


def can_use_source_docx_template(source_document):
    if Document is None:
        return False
    if not isinstance(source_document, dict):
        return False
    source_path = str(source_document.get("path") or "").strip()
    return bool(source_path and source_path.lower().endswith(".docx") and os.path.exists(source_path))


def add_docx_bullet_paragraph(document, text, bold=False):
    item = str(text or "").strip()
    if not item:
        return None
    try:
        paragraph = document.add_paragraph(style="List Bullet")
    except KeyError:
        paragraph = document.add_paragraph()
        paragraph.add_run("- ")
    run = paragraph.add_run(item)
    run.bold = bool(bold)
    return paragraph


def add_docx_bullet_list(document, items):
    for item in normalize_text_list(items):
        add_docx_bullet_paragraph(document, item)


def add_docx_text_section(document, title, lines):
    cleaned_lines = [str(line or "").strip() for line in lines or [] if str(line or "").strip()]
    if not cleaned_lines:
        return
    document.add_heading(title, level=2)
    for line in cleaned_lines:
        document.add_paragraph(line)


def iter_docx_paragraphs(document):
    """Yield paragraphs in the body and in table cells without rebuilding the document."""
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def change_replacements(change):
    """Turn scalar or structured change values into safe in-place text replacements."""
    change = change if isinstance(change, dict) else {}
    before = change.get("beforeValue")
    after = change.get("afterValue")
    replacements = []

    if isinstance(before, dict) and isinstance(after, dict):
        for key in ("title", "duration", "description", "note"):
            old = str(before.get(key) or "").strip()
            new = str(after.get(key) or "").strip()
            if old and new and old != new:
                replacements.append((old, new))
    elif isinstance(before, list) and isinstance(after, list):
        old_items = normalize_text_list(before)
        new_items = normalize_text_list(after)
        for old, new in zip(old_items, new_items):
            if old != new:
                replacements.append((old, new))
    else:
        old = str(before or "").strip()
        new = str(after or "").strip()
        if old and new and old not in ("n/a", "(not recorded)") and old != new:
            replacements.append((old, new))

    return replacements


def replace_in_docx_paragraph(paragraph, replacements):
    changed = 0
    has_drawing = "<w:drawing" in paragraph._p.xml or "<w:pict" in paragraph._p.xml
    for old, new in replacements:
        replaced_in_run = False
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                replaced_in_run = True
                changed += 1
        if not replaced_in_run and not has_drawing and old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new)
            changed += 1
    return changed


def change_location_key(change):
    change = change if isinstance(change, dict) else {}
    location = str(change.get("location") or "").strip()
    title = str(change.get("title") or "").strip()
    value = location.rsplit(">", 1)[-1].strip() if location else title
    value = re.sub(r"\s+(?:duration|content|note)?\s*updated$", "", value, flags=re.IGNORECASE)
    return normalize_heading_key(value)


def paragraphs_for_change(paragraphs, change):
    """Limit replacements to the located section or step when the change provides one."""
    target = change_location_key(change)
    if not target:
        return paragraphs

    matching_indexes = []
    for index, paragraph in enumerate(paragraphs):
        paragraph_key = normalize_heading_key(paragraph.text)
        if paragraph_key == target or target in paragraph_key:
            matching_indexes.append(index)

    if not matching_indexes:
        return []

    start = matching_indexes[0]
    end = len(paragraphs)
    for index in range(start + 1, len(paragraphs)):
        text = str(paragraphs[index].text or "").strip()
        if normalize_heading_key(text) in SECTION_HEADINGS or re.match(r"^\d+\s*[.)]\s+", text):
            end = index
            break
    return paragraphs[start:end]


def apply_tutor_changes_to_docx(document, final_plan):
    """Apply recorded changes at their original text locations and preserve the source layout."""
    changes = final_plan.get("changes") if isinstance(final_plan.get("changes"), list) else []
    paragraphs = list(iter_docx_paragraphs(document))
    applied = 0
    for change in changes:
        replacements = change_replacements(change)
        scoped_paragraphs = paragraphs_for_change(paragraphs, change)
        if not scoped_paragraphs:
            # A section may not have a recognizable heading in an imported document.
            # Only fall back when the source text is unambiguous across the document.
            scoped_paragraphs = [
                paragraph for paragraph in paragraphs
                if any(old in paragraph.text for old, _ in replacements)
            ]
            if sum(paragraph.text.count(old) for paragraph in scoped_paragraphs for old, _ in replacements) != 1:
                scoped_paragraphs = []
                log_message(f"[export] Could not safely locate DOCX change: {change.get('title', 'Change')}")
        for paragraph in scoped_paragraphs:
            applied += replace_in_docx_paragraph(paragraph, replacements)

    if changes and not applied:
        log_message("[export] No recorded DOCX change matched source text; original layout preserved")
    else:
        log_message(f"[export] Applied {applied} in-place DOCX text replacement(s)")


def render_source_docx_with_changes(context, output_path):
    final_plan = get_final_plan_context(context)
    source_document = final_plan.get("sourceDocument") if isinstance(final_plan.get("sourceDocument"), dict) else {}
    source_path = source_document.get("path", "")
    document = Document(source_path)
    apply_tutor_changes_to_docx(document, final_plan)
    document.save(output_path)
    log_message(f"[export] DOCX based on original template: {os.path.basename(source_path)}")
    return output_path


def render_lines_to_docx(lines, output_path):
    document = Document()
    lines = lines or []
    seen_title = False
    for line in lines:
        text = str(line or "")
        stripped = text.strip()
        if not stripped:
            document.add_paragraph("")
            continue
        if not seen_title:
            document.add_heading(stripped, level=0)
            seen_title = True
            continue
        normalized = normalize_heading_key(stripped)
        if normalized in SECTION_HEADINGS:
            document.add_heading(stripped, level=1)
            continue
        if stripped.startswith("- "):
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
            continue
        paragraph = document.add_paragraph()
        if re.match(r"^\d+\.\s", stripped):
            run = paragraph.add_run(stripped)
            run.bold = True
        else:
            paragraph.add_run(stripped)

    document.save(output_path)
    return output_path


def render_lines_to_pdf(lines, output_path):
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ExportTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1d2340"),
        spaceAfter=10,
        alignment=TA_LEFT,
    )
    heading_style = ParagraphStyle(
        "ExportHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#4f46e5"),
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "ExportBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#243044"),
        spaceAfter=4,
    )

    story = []
    seen_title = False
    for line in lines or []:
        text = str(line or "").strip()
        if not text:
            story.append(Spacer(1, 4))
            continue
        if not seen_title:
            story.append(Paragraph(xml_escape(text), title_style))
            seen_title = True
            continue
        normalized = normalize_heading_key(text)
        if normalized in SECTION_HEADINGS:
            story.append(Spacer(1, 6))
            story.append(Paragraph(xml_escape(text), heading_style))
            continue
        if text.startswith("- "):
            story.append(Paragraph(xml_escape("- " + text[2:].strip()), body_style))
            continue
        story.append(Paragraph(xml_escape(text), body_style))

    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )
    doc.build(story)
    return output_path


def ensure_export_folder(session_id):
    session_name = normalize_filename_piece(session_id or "session")
    export_folder = os.path.join(config.EXPORTS_ROOT, session_name)
    os.makedirs(export_folder, exist_ok=True)
    return export_folder


def build_docx_export(context, output_path):
    source_document = context.get("sourceDocument") if isinstance(context.get("sourceDocument"), dict) else {}
    if can_use_source_docx_template(source_document):
        try:
            return render_source_docx_with_changes(context, output_path)
        except Exception:
            log_message("[export] Template-based DOCX export failed; falling back to text-only DOCX")
            log_message(traceback.format_exc())

    lines, _ = build_final_export_lines(context)
    return render_lines_to_docx(lines, output_path)


def build_pdf_export(context, output_path):
    lines, _ = build_final_export_lines(context)
    return render_lines_to_pdf(lines, output_path)


def build_export_bundle(context):
    context = context if isinstance(context, dict) else {}
    session_id = str(context.get("sessionId", "")).strip() or "session"
    export_folder = ensure_export_folder(session_id)
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else {}
    source_document = context.get("sourceDocument") if isinstance(context.get("sourceDocument"), dict) else {}
    export_slug = normalize_filename_piece(draft.get("title") or source_document.get("filename") or "lesson-plan", max_length=72)
    docx_path = os.path.join(export_folder, f"{export_slug}.docx")
    pdf_path = os.path.join(export_folder, f"{export_slug}.pdf")

    build_docx_export(context, docx_path)
    build_pdf_export(context, pdf_path)

    return {
        "sessionId": session_id,
        "folder": export_folder,
        "docxPath": docx_path,
        "pdfPath": pdf_path,
        "filenameBase": export_slug,
    }

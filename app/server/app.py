import http.server
import os
import posixpath
import urllib.parse
import cgi
import json
import uuid
import webbrowser
import mimetypes
import zipfile
import tempfile
import urllib.request
import urllib.error
import traceback
import re
import unicodedata
import smtplib
import sys
from datetime import datetime
from io import BytesIO
import textwrap
from xml.etree import ElementTree as ET
from http import HTTPStatus
from email.message import EmailMessage
from xml.sax.saxutils import escape as xml_escape

RUNTIME_PYTHON_PACKAGES = r"C:\Users\Uinsel\.cache\codex-runtimes\codex-primary-runtime\dependencies\python"
if os.path.isdir(RUNTIME_PYTHON_PACKAGES) and RUNTIME_PYTHON_PACKAGES not in sys.path:
    sys.path.insert(0, RUNTIME_PYTHON_PACKAGES)

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
except ImportError:
    colors = None
    TA_LEFT = None
    A4 = None
    ParagraphStyle = None
    getSampleStyleSheet = None
    mm = None
    ListFlowable = None
    ListItem = None
    Paragraph = None
    SimpleDocTemplate = None
    Spacer = None

HERE = os.path.dirname(os.path.abspath(__file__))
APP_ROOT = os.path.abspath(os.path.join(HERE, ".."))
FRONTEND_ROOT = os.path.join(APP_ROOT, "frontend")
DATA_ROOT = os.path.join(APP_ROOT, "data")
PROMPTS_ROOT = os.path.join(HERE, "prompts")
CURRICULA_ROOT = os.path.join(DATA_ROOT, "curricula")
AMENDMENTS_ROOT = os.path.join(DATA_ROOT, "amendments")
LESSONPLANS_ROOT = os.path.join(DATA_ROOT, "lessonplans")
CONFIG_FILE = os.path.join(DATA_ROOT, "config.json")
ENV_FILE = os.path.join(HERE, ".env")
ANALYSIS_LOG_FILE = os.path.join(DATA_ROOT, "outputs", "analysis-log.txt")
MISTRAL_PROMPT_LOG_FILE = os.path.join(DATA_ROOT, "outputs", "mistral-prompt-log.txt")
STEP3_STATE_ROOT = os.path.join(DATA_ROOT, "outputs", "step3-sessions")
EXPORTS_ROOT = os.path.join(DATA_ROOT, "outputs", "exports")
MISTRAL_API_KEY = None
MISTRAL_MODEL = "mistral-small-latest"
SMTP_HOST = None
SMTP_PORT = 587
SMTP_USERNAME = None
SMTP_PASSWORD = None
SMTP_USE_TLS = True
SMTP_USE_SSL = False
MAIL_FROM_ADDRESS = None
MAIL_TO_ADDRESS = None

PORT = 8000

for path in (
    FRONTEND_ROOT,
    CURRICULA_ROOT,
    AMENDMENTS_ROOT,
    LESSONPLANS_ROOT,
    os.path.join(DATA_ROOT, "templates"),
    os.path.join(DATA_ROOT, "outputs"),
    STEP3_STATE_ROOT,
    EXPORTS_ROOT,
):
    os.makedirs(path, exist_ok=True)

if not os.path.exists(ANALYSIS_LOG_FILE):
    with open(ANALYSIS_LOG_FILE, "w", encoding="utf-8") as log_file:
        log_file.write("")

if not os.path.exists(MISTRAL_PROMPT_LOG_FILE):
    with open(MISTRAL_PROMPT_LOG_FILE, "w", encoding="utf-8") as log_file:
        log_file.write("")


def log_message(message):
    line = message.rstrip("\n")
    print(line)
    with open(ANALYSIS_LOG_FILE, "a", encoding="utf-8") as log_file:
        log_file.write(line + "\n")


def log_mistral_prompt(label, system_prompt, user_prompt, request_body=None):
    timestamp = datetime.now().isoformat(timespec="seconds")
    entry_lines = [
        "=" * 80,
        f"[{timestamp}] {label}",
        "",
        "SYSTEM PROMPT:",
        str(system_prompt or "").rstrip(),
        "",
        "USER PROMPT:",
        str(user_prompt or "").rstrip(),
    ]
    if request_body is not None:
        entry_lines.extend([
            "",
            "REQUEST BODY:",
            json.dumps(request_body, ensure_ascii=False, indent=2),
        ])
    entry_lines.append("")
    with open(MISTRAL_PROMPT_LOG_FILE, "a", encoding="utf-8") as log_file:
        log_file.write("\n".join(entry_lines) + "\n")


def load_prompt_file(filename):
    path = os.path.join(PROMPTS_ROOT, filename)
    try:
        with open(path, "r", encoding="utf-8") as prompt_file:
            return prompt_file.read().strip()
    except OSError as error:
        log_message(f"[prompt] Unable to load {filename}: {error}")
        raise


def load_env_file():
    global MISTRAL_API_KEY, MISTRAL_MODEL, SMTP_HOST, SMTP_PORT, SMTP_USERNAME, SMTP_PASSWORD, SMTP_USE_TLS, SMTP_USE_SSL, MAIL_FROM_ADDRESS, MAIL_TO_ADDRESS

    if os.path.exists(ENV_FILE):
        with open(ENV_FILE, "r", encoding="utf-8") as env_file:
            for raw_line in env_file:
                line = raw_line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                key, value = line.split("=", 1)
                value = value.strip().strip('"').strip("'")
                if key == "MISTRAL_API_KEY" and value:
                    MISTRAL_API_KEY = value
                elif key == "MISTRAL_MODEL" and value:
                    MISTRAL_MODEL = value
                elif key == "SMTP_HOST" and value:
                    SMTP_HOST = value
                elif key == "SMTP_PORT" and value:
                    try:
                        SMTP_PORT = int(value)
                    except ValueError:
                        SMTP_PORT = 587
                elif key == "SMTP_USERNAME" and value:
                    SMTP_USERNAME = value
                elif key == "SMTP_PASSWORD" and value:
                    SMTP_PASSWORD = value
                elif key == "SMTP_USE_TLS" and value:
                    SMTP_USE_TLS = value.lower() in ("1", "true", "yes", "on")
                elif key == "SMTP_USE_SSL" and value:
                    SMTP_USE_SSL = value.lower() in ("1", "true", "yes", "on")
                elif key == "MAIL_FROM_ADDRESS" and value:
                    MAIL_FROM_ADDRESS = value
                elif key == "MAIL_TO_ADDRESS" and value:
                    MAIL_TO_ADDRESS = value

    if not MISTRAL_API_KEY:
        MISTRAL_API_KEY = os.environ.get("MISTRAL_API_KEY")
    if os.environ.get("MISTRAL_MODEL"):
        MISTRAL_MODEL = os.environ["MISTRAL_MODEL"]
    if os.environ.get("SMTP_HOST"):
        SMTP_HOST = os.environ["SMTP_HOST"]
    if os.environ.get("SMTP_PORT"):
        try:
            SMTP_PORT = int(os.environ["SMTP_PORT"])
        except ValueError:
            SMTP_PORT = 587
    if os.environ.get("SMTP_USERNAME"):
        SMTP_USERNAME = os.environ["SMTP_USERNAME"]
    if os.environ.get("SMTP_PASSWORD"):
        SMTP_PASSWORD = os.environ["SMTP_PASSWORD"]
    if os.environ.get("SMTP_USE_TLS"):
        SMTP_USE_TLS = os.environ["SMTP_USE_TLS"].lower() in ("1", "true", "yes", "on")
    if os.environ.get("SMTP_USE_SSL"):
        SMTP_USE_SSL = os.environ["SMTP_USE_SSL"].lower() in ("1", "true", "yes", "on")
    if os.environ.get("MAIL_FROM_ADDRESS"):
        MAIL_FROM_ADDRESS = os.environ["MAIL_FROM_ADDRESS"]
    if os.environ.get("MAIL_TO_ADDRESS"):
        MAIL_TO_ADDRESS = os.environ["MAIL_TO_ADDRESS"]


def slugify_folder_name(value):
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-z0-9]+", "-", text)
    text = re.sub(r"-+", "-", text).strip("-")
    return text or "unknown"


def parse_subjects_value(value):
    if not value:
        return []
    if isinstance(value, list):
        values = value
    else:
        text = str(value).strip()
        if not text:
            return []
        try:
            parsed = json.loads(text)
            if isinstance(parsed, list):
                values = parsed
            else:
                values = [parsed]
        except Exception:
            values = [part.strip() for part in re.split(r"[;,]", text)]
    cleaned = []
    for item in values:
        item_text = str(item or "").strip()
        if item_text and item_text not in cleaned:
            cleaned.append(item_text)
    return cleaned


def ensure_step2_folder_structure(country, subjects):
    country_name = str(country or "").strip() or "unknown-country"
    country_slug = slugify_folder_name(country_name)
    created_paths = []

    curriculum_country_root = os.path.join(CURRICULA_ROOT, country_slug)
    amendments_country_root = os.path.join(AMENDMENTS_ROOT, country_slug)
    for root_path in (curriculum_country_root, amendments_country_root):
        os.makedirs(root_path, exist_ok=True)
        created_paths.append(root_path)

    normalized_subjects = parse_subjects_value(subjects)
    if not normalized_subjects:
        normalized_subjects = ["unknown-subject"]

    subject_paths = []
    for subject in normalized_subjects:
        subject_slug = slugify_folder_name(subject)
        curriculum_subject_path = os.path.join(curriculum_country_root, subject_slug)
        amendments_subject_path = os.path.join(amendments_country_root, subject_slug)
        os.makedirs(curriculum_subject_path, exist_ok=True)
        os.makedirs(amendments_subject_path, exist_ok=True)
        subject_paths.append(
            {
                "subject": subject,
                "slug": subject_slug,
                "curriculum": curriculum_subject_path,
                "amendments": amendments_subject_path,
            }
        )
        created_paths.extend([curriculum_subject_path, amendments_subject_path])

    return {
        "country": country_name,
        "countrySlug": country_slug,
        "curriculumRoot": curriculum_country_root,
        "amendmentsRoot": amendments_country_root,
        "subjects": subject_paths,
        "createdPaths": created_paths,
    }


def save_uploaded_files_to_folder(file_items, target_folder):
    saved_files = []
    os.makedirs(target_folder, exist_ok=True)

    for file_item in file_items:
        filename = os.path.basename(getattr(file_item, "filename", "") or "")
        if not filename:
            continue
        target_path = os.path.join(target_folder, filename)
        counter = 1
        base_name, extension = os.path.splitext(filename)
        while os.path.exists(target_path):
            target_path = os.path.join(target_folder, f"{base_name}-{counter}{extension}")
            counter += 1
        with open(target_path, "wb") as output_file:
            output_file.write(file_item.file.read())
        saved_files.append({
            "filename": os.path.basename(target_path),
            "path": target_path,
        })

    return saved_files


def step3_state_path(session_id):
    session = str(session_id or "").strip()
    if not session:
        return ""
    return os.path.join(STEP3_STATE_ROOT, f"{session}.json")


def step3_conversation_log_path(session_id):
    session = str(session_id or "").strip()
    if not session:
        return ""
    return os.path.join(STEP3_STATE_ROOT, f"{session}-conversation.txt")


def save_step3_state(session_id, state):
    path = step3_state_path(session_id)
    if not path:
        return False
    try:
        with open(path, "w", encoding="utf-8") as state_file:
            json.dump(state, state_file, ensure_ascii=False, indent=2)
        return True
    except OSError:
        log_message(f"[step3-state] Unable to save session={session_id}")
        traceback.print_exc()
        return False


def load_step3_state(session_id):
    path = step3_state_path(session_id)
    if not path or not os.path.exists(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as state_file:
            return json.load(state_file)
    except Exception:
        log_message(f"[step3-state] Unable to load session={session_id}")
        traceback.print_exc()
        return None


def parse_bool_value(value, default=False):
    if value is None:
        return default
    return str(value).strip().lower() in ("1", "true", "yes", "on")


def normalize_email_list(value):
    if not value:
        return []
    if isinstance(value, list):
        items = value
    else:
        items = re.split(r"[;,]", str(value))
    cleaned = []
    for item in items:
        text = str(item or "").strip()
        if text and text not in cleaned:
            cleaned.append(text)
    return cleaned


def build_draft_text(draft):
    draft = draft if isinstance(draft, dict) else {}
    lines = []
    lines.append(str(draft.get("title") or "Lesson Plan"))
    lines.append("")
    summary = str(draft.get("summary") or "").strip()
    if summary:
      lines.append("Summary")
      lines.append(summary)
      lines.append("")

    meta = draft.get("meta") if isinstance(draft.get("meta"), dict) else {}
    meta_lines = []
    for label, value in [
        ("Country", meta.get("country", "")),
        ("Subject(s)", ", ".join(meta.get("subjects", [])) if isinstance(meta.get("subjects"), list) else meta.get("subjects", "")),
        ("Grade range", meta.get("gradeRange", "")),
        ("Age range", meta.get("ageRange", "")),
        ("Focus", meta.get("focus", "")),
    ]:
        value_text = str(value or "").strip()
        if value_text:
            meta_lines.append(f"{label}: {value_text}")
    if meta_lines:
        lines.append("Meta")
        lines.extend(meta_lines)
        lines.append("")

    def add_list_section(title, items):
        items = items if isinstance(items, list) else []
        if not items:
            return
        lines.append(title)
        for item in items:
            lines.append(f"- {item}")
        lines.append("")

    add_list_section("Goals", draft.get("goals", []))
    add_list_section("Skills", draft.get("skills", []))

    steps = draft.get("steps", [])
    if isinstance(steps, list) and steps:
        lines.append("Steps")
        for index, step in enumerate(steps, start=1):
            step = step if isinstance(step, dict) else {}
            title = str(step.get("title", f"Step {index}")).strip()
            duration = str(step.get("duration", "")).strip()
            description = str(step.get("description", "")).strip()
            status = str(step.get("status", "")).strip()
            note = str(step.get("note", "")).strip()
            line = f"{index}. {title}" + (f" ({duration})" if duration else "")
            lines.append(line)
            if description:
                lines.append(f"   {description}")
            if status:
                lines.append(f"   Status: {status}")
            if note:
                lines.append(f"   Note: {note}")
        lines.append("")

    add_list_section("Materials", draft.get("materials", []))
    add_list_section("Assessment", draft.get("assessment", []))

    reflection = str(draft.get("reflection", "")).strip()
    if reflection:
        lines.append("Reflection")
        lines.append(reflection)
        lines.append("")

    changes = draft.get("changes", [])
    if isinstance(changes, list) and changes:
        lines.append("Change summary")
        for change in changes:
            change = change if isinstance(change, dict) else {}
            title = str(change.get("title", "Change")).strip()
            reason = str(change.get("reason", "")).strip()
            source = str(change.get("source", "")).strip()
            heading = title + (f" - {source}" if source else "")
            lines.append(heading)
            if reason:
                lines.append(reason)
        lines.append("")

    return "\n".join(lines).strip()


def build_conversation_text(conversation):
    conversation = conversation if isinstance(conversation, list) else []
    lines = []
    for message in conversation:
        message = message if isinstance(message, dict) else {}
        role = str(message.get("role", "tutor")).strip() or "tutor"
        text = str(message.get("text", "")).strip()
        time = str(message.get("time", "")).strip()
        prefix = f"[{time}] {role}: " if time else f"{role}: "
        lines.append(prefix + text)
    return "\n".join(lines).strip()


def get_step3_focus_label(payload):
    current_node = payload.get("currentNode") if isinstance(payload.get("currentNode"), dict) else {}
    focus_title = str(current_node.get("pointTitle", "")).strip()
    focus_label = str(current_node.get("label", "")).strip()
    focus_id = str(current_node.get("id", "")).strip()
    focus_parts = [item for item in (focus_label, focus_title, focus_id) if item]
    return " - ".join(focus_parts) if focus_parts else "No active focus"


def summarize_step3_changes(draft):
    draft = draft if isinstance(draft, dict) else {}
    changes = draft.get("changes") if isinstance(draft.get("changes"), list) else []
    lines = []
    for index, change in enumerate(changes[-8:], start=1):
        change = change if isinstance(change, dict) else {}
        title = str(change.get("title") or change.get("kind") or "Change").strip()
        reason = str(change.get("reason", "")).strip()
        before_value = str(change.get("beforeValue", "")).strip()
        after_value = str(change.get("afterValue", "")).strip()
        before_after = ""
        if before_value or after_value:
            before_after = f" ({before_value or '?'} -> {after_value or '?'})"
        line = f"{index}. {title}{before_after}"
        if reason:
            line += f": {reason}"
        lines.append(line)
    return lines


def build_step3_memory_summary(payload, conversation=None, draft=None):
    payload = payload if isinstance(payload, dict) else {}
    draft = draft if isinstance(draft, dict) else payload.get("draft")
    draft = draft if isinstance(draft, dict) else {}
    conversation = conversation if isinstance(conversation, list) else payload.get("conversation")
    conversation = conversation if isinstance(conversation, list) else []

    meta = payload.get("meta") if isinstance(payload.get("meta"), dict) else {}
    subjects = payload.get("subjects") if isinstance(payload.get("subjects"), list) else meta.get("subjects")
    subjects = subjects if isinstance(subjects, list) else []
    latest_teacher_reply = str(payload.get("userMessage", "")).strip()
    current_focus = get_step3_focus_label(payload)
    changes = summarize_step3_changes(draft)
    recent_conversation = build_conversation_text(conversation[-10:])

    lines = [
        "Step 3 local memory summary",
        f"Session: {payload.get('sessionId', '') or 'unknown'}",
        f"Country: {payload.get('country') or meta.get('country') or 'unknown'}",
        f"Subjects: {', '.join(str(item) for item in subjects) if subjects else 'unknown'}",
        f"Current focus: {current_focus}",
        f"Latest teacher reply: {latest_teacher_reply or '(none)'}",
        "",
        "Decisions and changes already applied:",
    ]
    lines.extend(changes if changes else ["(none yet)"])
    lines.extend([
        "",
        "Recent conversation:",
        recent_conversation or "(no conversation yet)",
        "",
        "Instruction for the next answer:",
        "Treat the recent conversation and applied changes as memory. Do not contradict them. "
        "If the teacher challenges a previous suggestion, acknowledge it and continue from the saved context.",
    ])
    return "\n".join(lines).strip()


def save_step3_conversation_log(session_id, conversation, payload, result):
    path = step3_conversation_log_path(session_id)
    if not path:
        return False
    try:
        draft = None
        if isinstance(result, dict) and isinstance(result.get("updated_draft"), dict):
            draft = result.get("updated_draft")
        elif isinstance(payload, dict) and isinstance(payload.get("draft"), dict):
            draft = payload.get("draft")

        memory_summary = build_step3_memory_summary(payload, conversation, draft)
        conversation_text = build_conversation_text(conversation)
        content = [
            memory_summary,
            "",
            "=" * 72,
            "Full conversation log",
            "=" * 72,
            conversation_text or "(no conversation yet)",
            "",
        ]
        with open(path, "w", encoding="utf-8") as log_file:
            log_file.write("\n".join(content))
        return True
    except OSError:
        log_message(f"[step3-conversation-log] Unable to save session={session_id}")
        traceback.print_exc()
        return False


def send_summary_email(payload):
    if not SMTP_HOST:
        return {"error": "SMTP_HOST is not configured in app/server/.env"}

    recipient = MAIL_TO_ADDRESS or str(payload.get("recipient", "")).strip()
    if not recipient:
        return {"error": "MAIL_TO_ADDRESS is not configured in app/server/.env"}

    recipients = normalize_email_list(recipient)
    if not recipients:
        return {"error": "No valid email recipient configured"}

    context = build_export_context(payload)
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else {}
    conversation = context.get("conversation") if isinstance(context.get("conversation"), list) else []
    session_id = str(context.get("sessionId", "")).strip()

    message = EmailMessage()
    message["Subject"] = f"STEaiM-CT lesson plan{f' ({session_id})' if session_id else ''}"
    message["From"] = MAIL_FROM_ADDRESS or SMTP_USERNAME or recipients[0]
    message["To"] = ", ".join(recipients)

    draft_text = build_draft_text(draft)
    conversation_text = build_conversation_text(conversation)

    body_lines = [
        "Here is the final lesson plan and the conversation from the local tutor.",
        "",
        "Lesson plan preview:",
        draft_text or "(no draft available)",
        "",
        "Conversation preview:",
        conversation_text or "(no conversation available)",
    ]
    message.set_content("\n".join(body_lines))

    try:
        bundle = build_export_bundle(context)
    except Exception:
        log_message("[mail] Export bundle could not be generated, sending text only")
        traceback.print_exc()
        bundle = None

    if conversation_text:
        message.add_attachment(conversation_text.encode("utf-8"), maintype="text", subtype="plain", filename="conversation.txt")
    if bundle and os.path.exists(bundle.get("pdfPath", "")):
        with open(bundle["pdfPath"], "rb") as pdf_file:
            pdf_bytes = pdf_file.read()
        message.add_attachment(pdf_bytes, maintype="application", subtype="pdf", filename=os.path.basename(bundle["pdfPath"]))
    elif draft_text:
        message.add_attachment(draft_text.encode("utf-8"), maintype="text", subtype="plain", filename="lesson-plan-draft.txt")

    try:
        if SMTP_USE_SSL:
            with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30) as smtp:
                if SMTP_USERNAME:
                    smtp.login(SMTP_USERNAME, SMTP_PASSWORD or "")
                smtp.send_message(message)
        else:
            with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=30) as smtp:
                if SMTP_USE_TLS:
                    smtp.starttls()
                if SMTP_USERNAME:
                    smtp.login(SMTP_USERNAME, SMTP_PASSWORD or "")
                smtp.send_message(message)
    except Exception as error:
        log_message(f"[mail] Send failed: {error}")
        traceback.print_exc()
        return {"error": f"Unable to send email: {error}"}

    return {
        "sent": True,
        "recipient": recipients,
        "subject": message["Subject"],
    }


def pdf_safe_text(value):
    text = str(value or "")
    replacements = {
        "–": "-",
        "—": "-",
        "•": "-",
        "“": "\"",
        "”": "\"",
        "‘": "'",
        "’": "'",
        "→": "->",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return text.encode("latin-1", "replace").decode("latin-1")


def pdf_escape(value):
    return pdf_safe_text(value).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf_from_draft(draft):
    draft = draft if isinstance(draft, dict) else {}
    page_width = 595.28
    page_height = 841.89
    margin_left = 48
    margin_right = 48
    margin_top = 54
    margin_bottom = 54
    usable_width = page_width - margin_left - margin_right

    blocks = []

    def add_block(text, size=11, leading=None, kind="body"):
        blocks.append({
            "text": text,
            "size": size,
            "leading": leading or int(size * 1.35),
            "kind": kind,
        })

    add_block("Lesson Plan Draft", 20, 24, "title")
    summary = str(draft.get("summary") or "").strip()
    if summary:
        add_block(summary, 11, 15)

    meta = draft.get("meta") if isinstance(draft.get("meta"), dict) else {}
    meta_lines = []
    meta_values = [
        ("Country", meta.get("country", "")),
        ("Subject(s)", ", ".join(meta.get("subjects", [])) if isinstance(meta.get("subjects"), list) else meta.get("subjects", "")),
        ("Grade range", meta.get("gradeRange", "")),
        ("Age range", meta.get("ageRange", "")),
    ]
    for label, value in meta_values:
        value_text = str(value or "").strip()
        if value_text:
            meta_lines.append(f"{label}: {value_text}")
    for line in meta_lines:
        add_block(line, 9.5, 12)

    def add_section(title, items):
        items = items if isinstance(items, list) else []
        if not items:
            return
        add_block(title, 13, 17, "heading")
        for item in items:
            add_block(f"- {item}", 10.5, 14)

    add_section("Goals", draft.get("goals", []))
    add_section("Skills", draft.get("skills", []))

    steps = draft.get("steps", [])
    if isinstance(steps, list) and steps:
        add_block("Steps", 13, 17, "heading")
        for index, step in enumerate(steps, start=1):
            step = step if isinstance(step, dict) else {}
            step_title = str(step.get("title", f"Step {index}")).strip()
            duration = str(step.get("duration", "")).strip()
            description = str(step.get("description", "")).strip()
            status = str(step.get("status", "")).strip()
            note = str(step.get("note", "")).strip()
            line = f"{index}. {step_title}" + (f" - {duration}" if duration else "")
            add_block(line, 11, 14)
            if description:
                add_block(description, 10, 13)
            if status:
                add_block(f"Status: {status}", 9.5, 12)
            if note:
                add_block(note, 9.5, 12)

    add_section("Materials", draft.get("materials", []))
    add_section("Assessment", draft.get("assessment", []))

    reflection = str(draft.get("reflection", "")).strip()
    if reflection:
        add_block("Reflection", 13, 17, "heading")
        add_block(reflection, 10.5, 14)

    changes = draft.get("changes", [])
    if isinstance(changes, list) and changes:
        add_block("Change summary", 13, 17, "heading")
        for change in changes:
            change = change if isinstance(change, dict) else {}
            change_title = str(change.get("title", "Change")).strip()
            change_reason = str(change.get("reason", "")).strip()
            change_source = str(change.get("source", "")).strip()
            header = change_title + (f" - {change_source}" if change_source else "")
            add_block(header, 10.5, 14)
            if change_reason:
                add_block(change_reason, 9.5, 12)
            before_value = str(change.get("beforeValue", "")).strip()
            after_value = str(change.get("afterValue", "")).strip()
            if before_value or after_value:
                add_block(f"Before: {before_value or 'n/a'}", 9.5, 12)
                add_block(f"After: {after_value or 'n/a'}", 9.5, 12)

    lines = []
    for block in blocks:
        text = pdf_safe_text(block["text"])
        wrapped = textwrap.wrap(text, width=max(20, int(usable_width / (block["size"] * 0.48))), break_long_words=False, break_on_hyphens=False) or [""]
        for wrapped_line in wrapped:
            lines.append({
                "text": wrapped_line,
                "size": block["size"],
                "leading": block["leading"],
                "kind": block["kind"],
            })

    pages = []
    current = []
    y = page_height - margin_top
    for line in lines:
        leading = line["leading"]
        if y - leading < margin_bottom:
            pages.append(current)
            current = []
            y = page_height - margin_top
        current.append({
            "x": margin_left,
            "y": y,
            "text": pdf_escape(line["text"]),
            "size": line["size"],
        })
        y -= leading
    if current:
        pages.append(current)
    if not pages:
        pages = [[]]

    def make_content_stream(page_lines):
        commands = []
        for line in page_lines:
            commands.append("BT")
            commands.append(f"/F1 {line['size']} Tf")
            commands.append(f"{line['x']} {line['y']} Td")
            commands.append(f"({line['text']}) Tj")
            commands.append("ET")
        return "\n".join(commands).encode("latin-1", "replace")

    objects = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")

    page_count = len(pages)
    page_object_numbers = [5 + index * 2 for index in range(page_count)]
    kids = " ".join(f"{number} 0 R" for number in page_object_numbers)
    objects.append(f"<< /Type /Pages /Count {page_count} /Kids [{kids}] >>".encode("ascii"))
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    for index, page_lines in enumerate(pages):
        content_object_number = 4 + index * 2
        content_stream = make_content_stream(page_lines)
        objects.append(f"<< /Length {len(content_stream)} >>\nstream\n".encode("ascii") + content_stream + b"\nendstream")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {page_width} {page_height}] /Resources << /Font << /F1 3 0 R >> >> /Contents {content_object_number} 0 R >>".encode("ascii")
        )

    pdf = BytesIO()
    pdf.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_number, obj in enumerate(objects, start=1):
        offsets.append(pdf.tell())
        pdf.write(f"{object_number} 0 obj\n".encode("ascii"))
        pdf.write(obj)
        pdf.write(b"\nendobj\n")

    xref_offset = pdf.tell()
    pdf.write(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    pdf.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.write(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.write(
        (
            "trailer\n"
            f"<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return pdf.getvalue()


def extract_text_from_file(file_path):
    lower_name = file_path.lower()
    if lower_name.endswith(".txt") or lower_name.endswith(".md") or lower_name.endswith(".csv"):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    if lower_name.endswith(".docx"):
        try:
            with zipfile.ZipFile(file_path) as docx_file:
                xml_data = docx_file.read("word/document.xml")
            root = ET.fromstring(xml_data)
            namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paragraphs = []
            for paragraph in root.findall(".//w:p", namespaces):
                parts = [node.text for node in paragraph.findall(".//w:t", namespaces) if node.text]
                text = "".join(parts).strip()
                if text:
                    paragraphs.append(text)
            extracted = "\n".join(paragraphs)
            log_message(f"[extract] DOCX extracted {len(extracted)} chars from {os.path.basename(file_path)}")
            return extracted
        except Exception:
            log_message(f"[extract] Failed to extract DOCX text from {os.path.basename(file_path)}")
            traceback.print_exc()
            return ""

    log_message(f"[extract] No extractor available for {os.path.basename(file_path)}")
    return ""


def normalize_filename_piece(value, default="lesson-plan"):
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-").lower()
    return text or default


def first_non_empty_value(*values):
    for value in values:
        if isinstance(value, dict):
            path_value = str(value.get("path", "")).strip()
            if path_value:
                return value
            stored_value = str(value.get("storedPath", "")).strip()
            if stored_value:
                return value
        elif isinstance(value, str) and value.strip():
            return value
    return None


def normalize_source_document(source_document):
    if isinstance(source_document, str):
        source_document = {
            "path": source_document,
        }
    if not isinstance(source_document, dict):
        return None

    normalized = dict(source_document)
    path = str(normalized.get("path", "") or normalized.get("storedPath", "") or normalized.get("filePath", "")).strip()
    if path:
        normalized["path"] = path
    filename = str(normalized.get("filename", "") or normalized.get("originalFilename", "") or os.path.basename(path)).strip()
    if filename:
        normalized["filename"] = filename
    stored_name = str(normalized.get("storedFilename", "")).strip()
    if stored_name:
        normalized["storedFilename"] = stored_name
    return normalized if path or filename else None


def find_uploaded_file_path(filename):
    filename = os.path.basename(str(filename or "").strip())
    if not filename:
        return ""

    direct_path = os.path.join(LESSONPLANS_ROOT, filename)
    if os.path.exists(direct_path):
        return direct_path

    matches = []
    for root, _, files in os.walk(LESSONPLANS_ROOT):
        if filename in files:
            candidate = os.path.join(root, filename)
            try:
                mtime = os.path.getmtime(candidate)
            except OSError:
                mtime = 0
            matches.append((mtime, candidate))

    if matches:
        matches.sort(key=lambda item: item[0], reverse=True)
        return matches[0][1]

    return ""


def resolve_source_document(context):
    context = context if isinstance(context, dict) else {}
    payload = context.get("payload") if isinstance(context.get("payload"), dict) else {}
    state = context.get("state") if isinstance(context.get("state"), dict) else {}

    candidates = [
        payload.get("sourceDocument"),
        payload.get("source_document"),
        payload.get("lessonPlanPath"),
        payload.get("lesson_plan_path"),
        state.get("sourceDocument"),
    ]
    if isinstance(state.get("payload"), dict):
        candidates.extend([
            state["payload"].get("sourceDocument"),
            state["payload"].get("source_document"),
            state["payload"].get("lessonPlanPath"),
            state["payload"].get("lesson_plan_path"),
        ])

    for candidate in candidates:
        normalized = normalize_source_document(candidate)
        if not normalized:
            continue
        path = normalized.get("path", "")
        if path and os.path.exists(path):
            return normalized

    filename_candidates = [
        payload.get("filename"),
        payload.get("sourceFilename"),
        payload.get("source_filename"),
        payload.get("lessonPlanFilename"),
        payload.get("lesson_plan_filename"),
    ]
    meta = payload.get("meta") if isinstance(payload.get("meta"), dict) else {}
    filename_candidates.extend([
        meta.get("filename"),
        meta.get("sourceFilename"),
        meta.get("source_filename"),
    ])
    if isinstance(state.get("payload"), dict):
        state_payload = state["payload"]
        filename_candidates.extend([
            state_payload.get("filename"),
            state_payload.get("sourceFilename"),
            state_payload.get("source_filename"),
        ])

    for filename in filename_candidates:
        found_path = find_uploaded_file_path(filename)
        if found_path:
            return {
                "path": found_path,
                "filename": os.path.basename(found_path),
            }

    return None


def split_text_to_paragraphs(text):
    blocks = []
    for raw_block in str(text or "").splitlines():
        line = raw_block.rstrip()
        if not line:
            blocks.append("")
            continue
        blocks.append(line)
    return blocks


def normalize_text_list(value):
    if not isinstance(value, list):
        return []
    cleaned = []
    for item in value:
        text = str(item or "").strip()
        if text:
            cleaned.append(text)
    return cleaned


def get_final_plan_context(context):
    context = context if isinstance(context, dict) else {}
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else {}
    source_document = context.get("sourceDocument") if isinstance(context.get("sourceDocument"), dict) else {}
    source_text = str(context.get("sourceText") or "").strip()
    title = str(draft.get("title") or source_document.get("filename") or "Lesson plan").strip()
    summary = str(draft.get("summary") or "").strip()
    meta = draft.get("meta") if isinstance(draft.get("meta"), dict) else {}
    goals = normalize_text_list(draft.get("goals"))
    skills = normalize_text_list(draft.get("skills"))
    materials = normalize_text_list(draft.get("materials"))
    assessment = normalize_text_list(draft.get("assessment"))
    steps = draft.get("steps") if isinstance(draft.get("steps"), list) else []
    reflection = str(draft.get("reflection") or "").strip()
    changes = collect_export_changes(context)
    conversation = context.get("conversation") if isinstance(context.get("conversation"), list) else []

    return {
        "title": title,
        "summary": summary,
        "meta": meta,
        "goals": goals,
        "skills": skills,
        "steps": steps,
        "materials": materials,
        "assessment": assessment,
        "reflection": reflection,
        "changes": changes,
        "conversation": conversation,
        "sourceDocument": source_document,
        "sourceText": source_text,
    }


def build_export_context(payload):
    payload = payload if isinstance(payload, dict) else {}
    session_id = str(payload.get("sessionId", "")).strip()
    state = load_step3_state(session_id) if session_id else None
    state = state if isinstance(state, dict) else {}
    state_payload = state.get("payload") if isinstance(state.get("payload"), dict) else {}

    draft = None
    if isinstance(state.get("draft"), dict):
        draft = state.get("draft")
    if not draft and isinstance(payload.get("draft"), dict):
        draft = payload.get("draft")
    if not draft and isinstance(state_payload.get("draft"), dict):
        draft = state_payload.get("draft")
    if not isinstance(draft, dict):
        draft = {}

    conversation = None
    if isinstance(state.get("conversation"), list):
        conversation = state.get("conversation")
    if conversation is None and isinstance(payload.get("conversation"), list):
        conversation = payload.get("conversation")
    if conversation is None and isinstance(state_payload.get("conversation"), list):
        conversation = state_payload.get("conversation")
    if not isinstance(conversation, list):
        conversation = []

    source_document = resolve_source_document({
        "payload": payload,
        "state": state,
    })

    if not source_document and isinstance(state_payload, dict):
        source_document = resolve_source_document({
            "payload": state_payload,
            "state": state,
        })

    source_text = ""
    if source_document and source_document.get("path") and os.path.exists(source_document["path"]):
        source_text = extract_text_from_file(source_document["path"])

    if not source_text:
        source_text = str(draft.get("summary") or "").strip()

    if not source_document and source_text:
        source_document = {
            "filename": "lesson-plan",
            "path": "",
        }

    return {
        "sessionId": session_id,
        "state": state,
        "payload": payload,
        "draft": draft,
        "conversation": conversation,
        "sourceDocument": source_document,
        "sourceText": source_text,
    }


def collect_export_changes(context):
    context = context if isinstance(context, dict) else {}
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else {}
    state = context.get("state") if isinstance(context.get("state"), dict) else {}
    payload = context.get("payload") if isinstance(context.get("payload"), dict) else {}
    result = state.get("result") if isinstance(state.get("result"), dict) else {}

    raw_changes = []
    if isinstance(draft.get("changes"), list) and draft.get("changes"):
        raw_changes = draft.get("changes")
    elif isinstance(result.get("changes_made"), list) and result.get("changes_made"):
        raw_changes = result.get("changes_made")
    elif isinstance(payload.get("changes_made"), list) and payload.get("changes_made"):
        raw_changes = payload.get("changes_made")

    changes = []
    for item in raw_changes:
        item = item if isinstance(item, dict) else {}
        title = str(item.get("title") or item.get("section") or item.get("kind") or "Change").strip()
        reason = str(item.get("reason") or item.get("summary") or item.get("note") or "").strip()
        before_value = str(item.get("beforeValue") or item.get("before") or "").strip()
        after_value = str(item.get("afterValue") or item.get("after") or "").strip()
        source = str(item.get("source") or item.get("section") or "").strip()
        changes.append({
            "title": title,
            "reason": reason,
            "beforeValue": before_value,
            "afterValue": after_value,
            "source": source,
        })

    return changes


def build_export_metadata_lines(context):
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else {}
    meta = draft.get("meta") if isinstance(draft.get("meta"), dict) else {}
    source_document = context.get("sourceDocument") if isinstance(context.get("sourceDocument"), dict) else {}
    lines = []
    for label, value in [
        ("Country", meta.get("country", "")),
        ("Subject(s)", ", ".join(meta.get("subjects", [])) if isinstance(meta.get("subjects"), list) else meta.get("subjects", "")),
        ("Grade range", meta.get("gradeRange", "")),
        ("Age range", meta.get("ageRange", "")),
        ("Focus", meta.get("focus", "")),
        ("Original file", source_document.get("filename", "")),
    ]:
        value_text = str(value or "").strip()
        if value_text:
            lines.append(f"{label}: {value_text}")
    return lines


SECTION_HEADINGS = ("summary", "goals", "skills", "steps", "materials", "assessment", "reflection")


def normalize_heading_key(value):
    text = str(value or "").strip().lower()
    text = re.sub(r"^\s*\d+\s*[\.\)]\s*", "", text)
    text = text.replace("subject(s)", "subjects")
    text = text.replace("/", " ")
    text = re.sub(r"[^a-z0-9]+", " ", text).strip()
    return text


def build_section_lines(final_plan):
    lines = []

    summary = str(final_plan.get("summary") or "").strip()
    if summary:
        lines.append(summary)
        lines.append("")

    meta = final_plan.get("meta") if isinstance(final_plan.get("meta"), dict) else {}
    meta_lines = []
    for label, value in [
        ("Country", meta.get("country", "")),
        ("Subject(s)", ", ".join(meta.get("subjects", [])) if isinstance(meta.get("subjects"), list) else meta.get("subjects", "")),
        ("Grade range", meta.get("gradeRange", "")),
        ("Age range", meta.get("ageRange", "")),
        ("Focus", meta.get("focus", "")),
    ]:
        value_text = str(value or "").strip()
        if value_text:
            meta_lines.append(f"{label}: {value_text}")
    if meta_lines:
        lines.extend(meta_lines)
        lines.append("")

    def add_bullet_section(title, items):
        items = normalize_text_list(items)
        if not items:
            return
        lines.append(title)
        for item in items:
            lines.append(f"- {item}")
        lines.append("")

    add_bullet_section("Goals", final_plan.get("goals"))
    add_bullet_section("Skills", final_plan.get("skills"))

    steps = final_plan.get("steps") if isinstance(final_plan.get("steps"), list) else []
    if steps:
        lines.append("Steps")
        for index, step in enumerate(steps, start=1):
            step = step if isinstance(step, dict) else {}
            title = str(step.get("title") or f"Step {index}").strip()
            duration = str(step.get("duration") or "").strip()
            description = str(step.get("description") or "").strip()
            status = str(step.get("status") or "").strip()
            note = str(step.get("note") or "").strip()
            header = f"{index}. {title}" + (f" ({duration})" if duration else "")
            lines.append(header)
            if description:
                lines.append(description)
            if status:
                lines.append(f"Status: {status}")
            if note:
                lines.append(f"Note: {note}")
        lines.append("")

    add_bullet_section("Materials", final_plan.get("materials"))
    add_bullet_section("Assessment", final_plan.get("assessment"))

    reflection = str(final_plan.get("reflection") or "").strip()
    if reflection:
        lines.append("Reflection")
        lines.append(reflection)
        lines.append("")

    return lines


def replace_first_non_empty_line(lines, new_value):
    if not lines:
        return [new_value] if new_value else []
    updated = list(lines)
    for index, line in enumerate(updated):
        if str(line or "").strip():
            if new_value:
                updated[index] = new_value
            return updated
    if new_value:
        updated.insert(0, new_value)
    return updated


def replace_metadata_lines(lines, metadata_lines):
    updated = list(lines)
    for meta_line in metadata_lines:
        label = meta_line.split(":", 1)[0].strip().lower() if ":" in meta_line else ""
        if not label:
            continue
        replaced = False
        for index, line in enumerate(updated):
            normalized = str(line or "").strip().lower()
            if normalized.startswith(label + ":") or normalized == label:
                updated[index] = meta_line
                replaced = True
                break
        if not replaced:
            insert_at = 1 if updated else 0
            while insert_at < len(updated) and not str(updated[insert_at]).strip():
                insert_at += 1
            updated.insert(insert_at, meta_line)
            updated.insert(insert_at + 1, "")
    return updated


def apply_inline_replacements(lines, changes):
    updated = list(lines)
    changes = changes if isinstance(changes, list) else []
    for change in changes:
        before_value = str((change or {}).get("beforeValue") or "").strip()
        after_value = str((change or {}).get("afterValue") or "").strip()
        if not before_value or not after_value:
            continue
        updated = [line.replace(before_value, after_value) for line in updated]
    return updated


def replace_section_block(lines, heading, content_lines):
    content_lines = list(content_lines or [])
    normalized_heading = normalize_heading_key(heading)
    updated = list(lines)
    heading_index = None

    for index, line in enumerate(updated):
        if normalize_heading_key(line) == normalized_heading:
            heading_index = index
            break

    if heading_index is None:
        if content_lines:
            if updated and str(updated[-1]).strip():
                updated.append("")
            updated.append(heading)
            updated.extend(content_lines)
        return updated

    end_index = len(updated)
    for index in range(heading_index + 1, len(updated)):
        candidate = updated[index]
        candidate_key = normalize_heading_key(candidate)
        if candidate_key in SECTION_HEADINGS and candidate_key != normalized_heading:
            end_index = index
            break

    updated = updated[: heading_index + 1] + content_lines + updated[end_index:]
    return updated


def build_final_export_lines(context):
    final_plan = get_final_plan_context(context)
    source_text = str(final_plan.get("sourceText") or "").strip()
    lines = split_text_to_paragraphs(source_text) if source_text else []
    lines = apply_inline_replacements(lines, final_plan.get("changes"))
    lines = replace_first_non_empty_line(lines, final_plan["title"])
    lines = replace_metadata_lines(lines, build_export_metadata_lines(context))
    lines = replace_section_block(lines, "Summary", [final_plan["summary"]] if final_plan["summary"] else [])
    lines = replace_section_block(lines, "Goals", [f"- {item}" for item in normalize_text_list(final_plan.get("goals"))])
    lines = replace_section_block(lines, "Skills", [f"- {item}" for item in normalize_text_list(final_plan.get("skills"))])

    step_lines = []
    steps = final_plan.get("steps") if isinstance(final_plan.get("steps"), list) else []
    for index, step in enumerate(steps, start=1):
        step = step if isinstance(step, dict) else {}
        title = str(step.get("title") or f"Step {index}").strip()
        duration = str(step.get("duration") or "").strip()
        description = str(step.get("description") or "").strip()
        status = str(step.get("status") or "").strip()
        note = str(step.get("note") or "").strip()
        step_lines.append(f"{index}. {title}" + (f" ({duration})" if duration else ""))
        if description:
            step_lines.append(description)
        if status:
            step_lines.append(f"Status: {status}")
        if note:
            step_lines.append(f"Note: {note}")
        step_lines.append("")
    if step_lines and step_lines[-1] == "":
        step_lines.pop()
    lines = replace_section_block(lines, "Steps", step_lines)
    lines = replace_section_block(lines, "Materials", [f"- {item}" for item in normalize_text_list(final_plan.get("materials"))])
    lines = replace_section_block(lines, "Assessment", [f"- {item}" for item in normalize_text_list(final_plan.get("assessment"))])
    lines = replace_section_block(lines, "Reflection", [final_plan["reflection"]] if final_plan["reflection"] else [])

    lines = [line for line in lines if line is not None]
    while lines and not str(lines[0]).strip():
        lines.pop(0)
    return lines, final_plan


def render_lines_to_docx(lines, output_path):
    document = Document()
    lines = lines or []
    seen_title = False
    for line in lines:
        text = str(line or "")
        stripped = text.strip()
        if not stripped:
            document.add_paragraph("")
            continue
        if not seen_title:
            document.add_heading(stripped, level=0)
            seen_title = True
            continue
        normalized = normalize_heading_key(stripped)
        if normalized in SECTION_HEADINGS:
            document.add_heading(stripped, level=1)
            continue
        if stripped.startswith("- "):
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
            continue
        paragraph = document.add_paragraph()
        if re.match(r"^\d+\.\s", stripped):
            run = paragraph.add_run(stripped)
            run.bold = True
        else:
            paragraph.add_run(stripped)

    document.save(output_path)
    return output_path


def render_lines_to_pdf(lines, output_path):
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ExportTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1d2340"),
        spaceAfter=10,
        alignment=TA_LEFT,
    )
    heading_style = ParagraphStyle(
        "ExportHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#4f46e5"),
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "ExportBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#243044"),
        spaceAfter=4,
    )

    story = []
    seen_title = False
    for line in lines or []:
        text = str(line or "").strip()
        if not text:
            story.append(Spacer(1, 4))
            continue
        if not seen_title:
            story.append(Paragraph(xml_escape(text), title_style))
            seen_title = True
            continue
        normalized = normalize_heading_key(text)
        if normalized in SECTION_HEADINGS:
            story.append(Spacer(1, 6))
            story.append(Paragraph(xml_escape(text), heading_style))
            continue
        if text.startswith("- "):
            story.append(Paragraph(xml_escape("• " + text[2:].strip()), body_style))
            continue
        story.append(Paragraph(xml_escape(text), body_style))

    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )
    doc.build(story)
    return output_path


def ensure_export_folder(session_id):
    session_name = normalize_filename_piece(session_id or "session")
    export_folder = os.path.join(EXPORTS_ROOT, session_name)
    os.makedirs(export_folder, exist_ok=True)
    return export_folder


def build_docx_export(context, output_path):
    lines, _ = build_final_export_lines(context)
    return render_lines_to_docx(lines, output_path)


def build_pdf_export(context, output_path):
    lines, _ = build_final_export_lines(context)
    return render_lines_to_pdf(lines, output_path)


def build_export_bundle(context):
    context = context if isinstance(context, dict) else {}
    session_id = str(context.get("sessionId", "")).strip() or "session"
    export_folder = ensure_export_folder(session_id)
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else {}
    source_document = context.get("sourceDocument") if isinstance(context.get("sourceDocument"), dict) else {}
    export_slug = normalize_filename_piece(draft.get("title") or source_document.get("filename") or "lesson-plan")
    docx_path = os.path.join(export_folder, f"{export_slug}.docx")
    pdf_path = os.path.join(export_folder, f"{export_slug}.pdf")

    build_docx_export(context, docx_path)
    build_pdf_export(context, pdf_path)

    return {
        "sessionId": session_id,
        "folder": export_folder,
        "docxPath": docx_path,
        "pdfPath": pdf_path,
        "filenameBase": export_slug,
    }


def extract_json_block(text):
    if not text:
        return ""

    cleaned = text.strip()
    if cleaned.startswith("```"):
        lines = cleaned.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip().startswith("```"):
            lines = lines[:-1]
        cleaned = "\n".join(lines).strip()

    start = cleaned.find("{")
    if start == -1:
        return ""

    in_string = False
    escape = False
    depth = 0
    end = None

    for index in range(start, len(cleaned)):
        char = cleaned[index]
        if escape:
            escape = False
            continue
        if char == "\\":
            escape = True
            continue
        if char == '"':
            in_string = not in_string
            continue
        if in_string:
            continue
        if char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                end = index + 1
                break

    if end is None:
        return cleaned[start:].strip()
    return cleaned[start:end].strip()


def repair_json_quotes(text):
    if not text:
        return text

    repaired = []
    in_string = False
    escape = False

    for index, char in enumerate(text):
        if escape:
            repaired.append(char)
            escape = False
            continue

        if char == "\\":
            repaired.append(char)
            if in_string:
                escape = True
            continue

        if char == '"':
            if not in_string:
                in_string = True
                repaired.append(char)
                continue

            next_index = index + 1
            while next_index < len(text) and text[next_index].isspace():
                next_index += 1
            next_char = text[next_index] if next_index < len(text) else ""

            if next_char in (",", "}", "]", "", ":"):
                in_string = False
                repaired.append(char)
            else:
                repaired.append("\\\"")
            continue

        repaired.append(char)

    return "".join(repaired)


def parse_json_response_block(text):
    raw_block = extract_json_block(text)
    if not raw_block:
        return None, None

    try:
        return json.loads(raw_block), raw_block
    except json.JSONDecodeError as first_error:
        repaired = repair_json_quotes(raw_block)
        if repaired != raw_block:
            try:
                return json.loads(repaired), repaired
            except json.JSONDecodeError as second_error:
                raise second_error from first_error
        raise first_error


def format_context_range(from_value, to_value, unknown_label="not provided"):
    start = str(from_value or "").strip()
    end = str(to_value or "").strip()
    if start and end:
        return f"{start} to {end}"
    if start:
        return f"{start} to ?"
    if end:
        return f"? to {end}"
    return unknown_label


def build_ui_target_context(metadata):
    subjects = metadata.get("subjects", [])
    if isinstance(subjects, str):
        subjects = parse_subjects_value(subjects)
    return {
        "country": str(metadata.get("country", "")).strip(),
        "subjects": subjects,
        "gradeRange": {
            "from": str(metadata.get("gradeFrom", "")).strip(),
            "to": str(metadata.get("gradeTo", "")).strip(),
            "display": format_context_range(metadata.get("gradeFrom", ""), metadata.get("gradeTo", "")),
        },
        "ageRange": {
            "from": str(metadata.get("ageFrom", "")).strip(),
            "to": str(metadata.get("ageTo", "")).strip(),
            "display": format_context_range(metadata.get("ageFrom", ""), metadata.get("ageTo", "")),
        },
        "specifics": str(metadata.get("specifics", "")).strip(),
        "priority": "Use this teacher-entered UI context as the intended lesson context. Check the uploaded document against it.",
    }


def build_analysis_context_summary(payload):
    context = payload.get("ui_target_context") if isinstance(payload.get("ui_target_context"), dict) else {}
    subjects = context.get("subjects") if isinstance(context.get("subjects"), list) else []
    grade_range = context.get("gradeRange") if isinstance(context.get("gradeRange"), dict) else {}
    age_range = context.get("ageRange") if isinstance(context.get("ageRange"), dict) else {}
    lines = [
        "Teacher-entered UI target context (read this first):",
        f"- Country: {context.get('country') or 'not provided'}",
        f"- Subject(s): {', '.join(subjects) if subjects else 'not provided'}",
        f"- Intended grade range: {grade_range.get('display') or 'not provided'}",
        f"- Intended age range: {age_range.get('display') or 'not provided'}",
        f"- Specific teacher notes: {context.get('specifics') or 'none'}",
        "",
        "Analysis instruction:",
        "Use the UI target context as the teacher's intended use case. Extract what the document says, then check whether the lesson plan fits this UI context based on the actual tasks and activity load. If the document claims a different grade or age than the UI context, report the mismatch clearly.",
    ]
    return "\n".join(lines)


def call_mistral_analysis(payload):
    if not MISTRAL_API_KEY:
        log_message("[mistral] Missing MISTRAL_API_KEY")
        return {
            "error": "MISTRAL_API_KEY is not configured in app/server/.env"
        }

    system_prompt = load_prompt_file("analysis_system_prompt.txt")

    context_summary = build_analysis_context_summary(payload)
    user_prompt = context_summary + "\n\nFull uploaded lesson payload:\n" + json.dumps(payload, ensure_ascii=False, indent=2)
    request_body = {
        "model": MISTRAL_MODEL,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.2,
    }
    log_mistral_prompt("analysis", system_prompt, user_prompt, request_body)

    request = urllib.request.Request(
        "https://api.mistral.ai/v1/chat/completions",
        data=json.dumps(request_body).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {MISTRAL_API_KEY}",
            "Content-Type": "application/json",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            response_data = json.loads(response.read().decode("utf-8"))
            log_message(f"[mistral] Response received, model={MISTRAL_MODEL}")
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="ignore")
        log_message(f"[mistral] HTTPError {error.code}: {detail}")
        return {"error": f"Mistral request failed with HTTP {error.code}", "detail": detail}
    except Exception as error:
        log_message(f"[mistral] Request failed: {error}")
        traceback.print_exc()
        return {"error": f"Mistral request failed: {error}"}

    try:
        content = response_data["choices"][0]["message"]["content"]
        log_message(f"[mistral] Raw content preview: {content[:500]}")
        parsed, json_block = parse_json_response_block(content)
        if not json_block:
            return {"error": "Mistral did not return a JSON block", "raw": content}
        if isinstance(parsed, dict):
            return parsed
        return {"error": "Mistral did not return a JSON object", "raw": json_block}
    except Exception as error:
        log_message(f"[mistral] Could not parse response: {error}")
        traceback.print_exc()
        return {"error": f"Could not parse Mistral response as JSON: {error}", "raw": response_data}


def call_mistral_refinement(payload):
    if not MISTRAL_API_KEY:
        log_message("[mistral-step3] Missing MISTRAL_API_KEY")
        return {
            "error": "MISTRAL_API_KEY is not configured in app/server/.env"
        }

    system_prompt = load_prompt_file("refinement_system_prompt.txt")

    memory_summary = build_step3_memory_summary(payload)
    raw_payload = json.dumps(payload, ensure_ascii=False, indent=2)
    user_prompt = (
        memory_summary
        + "\n\n"
        + "=" * 72
        + "\nRaw current payload follows. Use it as source data, but prioritize the memory summary for dialogue continuity.\n"
        + "=" * 72
        + "\n"
        + raw_payload
    )
    request_body = {
        "model": MISTRAL_MODEL,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.2,
    }
    log_mistral_prompt("refinement", system_prompt, user_prompt, request_body)

    request = urllib.request.Request(
        "https://api.mistral.ai/v1/chat/completions",
        data=json.dumps(request_body).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {MISTRAL_API_KEY}",
            "Content-Type": "application/json",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            response_data = json.loads(response.read().decode("utf-8"))
            log_message(f"[mistral-step3] Response received, model={MISTRAL_MODEL}")
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="ignore")
        log_message(f"[mistral-step3] HTTPError {error.code}: {detail}")
        return {"error": f"Mistral refinement failed with HTTP {error.code}", "detail": detail}
    except Exception as error:
        log_message(f"[mistral-step3] Request failed: {error}")
        traceback.print_exc()
        return {"error": f"Mistral refinement failed: {error}"}

    try:
        content = response_data["choices"][0]["message"]["content"]
        log_message(f"[mistral-step3] Raw content preview: {content[:500]}")
        parsed, json_block = parse_json_response_block(content)
        if not json_block:
            return {"error": "Mistral did not return a JSON block", "raw": content}
        if isinstance(parsed, dict):
            return parsed
        return {"error": "Mistral did not return a JSON object", "raw": json_block}
    except Exception as error:
        log_message(f"[mistral-step3] Could not parse response: {error}")
        traceback.print_exc()
        return {"error": f"Could not parse Mistral step 3 response as JSON: {error}", "raw": response_data}


def call_mistral_hello():
    if not MISTRAL_API_KEY:
        log_message("[mistral-test] Missing MISTRAL_API_KEY")
        return {"error": "MISTRAL_API_KEY is not configured in app/server/.env"}

    system_prompt = load_prompt_file("mistral_test_system_prompt.txt")
    request_body = {
        "model": MISTRAL_MODEL,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": "Say hello."},
        ],
        "temperature": 0,
    }
    log_mistral_prompt("test", system_prompt, "Say hello.", request_body)

    request = urllib.request.Request(
        "https://api.mistral.ai/v1/chat/completions",
        data=json.dumps(request_body).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {MISTRAL_API_KEY}",
            "Content-Type": "application/json",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            response_data = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="ignore")
        log_message(f"[mistral-test] HTTPError {error.code}: {detail}")
        return {"error": f"Mistral test failed with HTTP {error.code}", "detail": detail}
    except Exception as error:
        log_message(f"[mistral-test] Request failed: {error}")
        traceback.print_exc()
        return {"error": f"Mistral test failed: {error}"}

    try:
        content = response_data["choices"][0]["message"]["content"].strip().lower()
        log_message(f"[mistral-test] Raw content: {content}")
        if "hello" in content:
            return {"hello": "hello", "raw": content}
        return {"hello": content, "raw": content}
    except Exception as error:
        log_message(f"[mistral-test] Could not parse response: {error}")
        traceback.print_exc()
        return {"error": f"Could not parse Mistral test response: {error}", "raw": response_data}


load_env_file()

if not os.path.exists(CONFIG_FILE):
    default_config = {
        "countries": [
            "Slovakia",
            "Germany",
            "Austria",
            "Spain",
            "Czech Republic",
            "Poland",
            "Portugal",
            "Italy"
        ],
        "subjects": [
            "Mathematics",
            "Informatics / CS",
            "Biology",
            "Physics",
            "Chemistry"
        ]
    }
    with open(CONFIG_FILE, "w", encoding="utf-8") as config_file:
        json.dump(default_config, config_file, indent=2)

class TutorHandler(http.server.BaseHTTPRequestHandler):
    def do_GET(self):
        route = urllib.parse.urlparse(self.path).path

        if route in ("", "/"):
            self.serve_file(os.path.join(FRONTEND_ROOT, "start.html"), "text/html")
        elif route == "/ping":
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(json.dumps({"status": "ok"}).encode("utf-8"))
        elif route == "/step3-state":
            self.handle_step3_state()
        elif route == "/config":
            self.serve_config()
        elif route.startswith("/data/curricula/"):
            self.serve_path(route[len("/data/curricula/"):], CURRICULA_ROOT)
        elif route.startswith("/frontend/"):
            self.serve_path(route[len("/frontend/"):], FRONTEND_ROOT)
        else:
            self.serve_path(route.lstrip("/"), FRONTEND_ROOT)

    def do_POST(self):
        route = urllib.parse.urlparse(self.path).path
        if route == "/upload":
            self.handle_upload()
        elif route == "/analyze":
            self.handle_analyze()
        elif route == "/step2-prepare":
            self.handle_step2_prepare()
        elif route == "/step3-refine":
            self.handle_step3_refine()
        elif route == "/send-email":
            self.handle_send_email()
        elif route == "/export-docx":
            self.handle_export_docx()
        elif route == "/export-pdf":
            self.handle_export_pdf()
        elif route == "/mistral-test":
            self.handle_mistral_test()
        elif route == "/client-log":
            self.handle_client_log()
        elif route == "/config":
            self.handle_config_update()
        else:
            self.send_error(HTTPStatus.NOT_FOUND, "Not found")

    def handle_upload(self):
        content_type = self.headers.get("Content-Type", "")
        if not content_type.startswith("multipart/form-data"):
            self.send_error(HTTPStatus.BAD_REQUEST, "Expected multipart/form-data")
            return

        form = cgi.FieldStorage(
            fp=self.rfile,
            headers=self.headers,
            environ={
                "REQUEST_METHOD": "POST",
                "CONTENT_TYPE": content_type,
            },
        )

        if "lessonPlan" not in form and "file" not in form:
            self.send_error(HTTPStatus.BAD_REQUEST, "No file field provided")
            return

        file_item = form["lessonPlan"] if "lessonPlan" in form else form["file"]
        filename = os.path.basename(file_item.filename or "")

        if not filename:
            self.send_error(HTTPStatus.BAD_REQUEST, "Invalid filename")
            return

        target_path = os.path.join(LESSONPLANS_ROOT, filename)
        try:
            with open(target_path, "wb") as output_file:
                output_file.write(file_item.file.read())
        except OSError:
            self.send_error(HTTPStatus.INTERNAL_SERVER_ERROR, "Unable to save file")
            return

        specifics = form.getvalue("specifics", "").strip()
        specifics_filename = None
        if specifics:
            specifics_filename = f"specifics-{uuid.uuid4().hex[:8]}.txt"
            specifics_path = os.path.join(DATA_ROOT, "outputs", specifics_filename)
            try:
                with open(specifics_path, "w", encoding="utf-8") as specifics_file:
                    specifics_file.write(specifics)
            except OSError:
                self.send_error(HTTPStatus.INTERNAL_SERVER_ERROR, "Unable to save specifics")
                return

        response_body = {"saved": filename}
        if specifics_filename:
            response_body["specificsSaved"] = specifics_filename

        self.send_response(HTTPStatus.CREATED)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps(response_body).encode("utf-8"))

    def handle_analyze(self):
        content_type = self.headers.get("Content-Type", "")
        if not content_type.startswith("multipart/form-data"):
            self.send_error(HTTPStatus.BAD_REQUEST, "Expected multipart/form-data")
            return

        form = cgi.FieldStorage(
            fp=self.rfile,
            headers=self.headers,
            environ={
                "REQUEST_METHOD": "POST",
                "CONTENT_TYPE": content_type,
            },
        )

        if "lessonPlan" not in form and "file" not in form:
            self.send_error(HTTPStatus.BAD_REQUEST, "No file field provided")
            return

        file_item = form["lessonPlan"] if "lessonPlan" in form else form["file"]
        filename = os.path.basename(file_item.filename or "")
        if not filename:
            self.send_error(HTTPStatus.BAD_REQUEST, "Invalid filename")
            return

        log_message(f"[analyze] Received file={filename}")

        temp_suffix = os.path.splitext(filename)[1] or ".bin"
        file_bytes = file_item.file.read()
        with tempfile.NamedTemporaryFile(delete=False, suffix=temp_suffix) as temp_file:
            temp_file.write(file_bytes)
            temp_path = temp_file.name

        try:
            extracted_text = extract_text_from_file(temp_path)
        finally:
                try:
                    os.remove(temp_path)
                except OSError:
                    pass

        lessonplan_upload_root = os.path.join(LESSONPLANS_ROOT, "uploads")
        os.makedirs(lessonplan_upload_root, exist_ok=True)
        stored_filename = f"{uuid.uuid4().hex[:12]}-{normalize_filename_piece(os.path.splitext(filename)[0])}{temp_suffix.lower()}"
        stored_path = os.path.join(lessonplan_upload_root, stored_filename)
        try:
            with open(stored_path, "wb") as output_file:
                output_file.write(file_bytes)
        except OSError:
            self.send_error(HTTPStatus.INTERNAL_SERVER_ERROR, "Unable to save uploaded lesson plan")
            return

        log_message(f"[analyze] Extracted chars={len(extracted_text)}")

        specifics = form.getvalue("specifics", "").strip()
        subjects_value = []
        if "subjects[]" in form:
            subjects_field = form["subjects[]"]
            if isinstance(subjects_field, list):
                subjects_value = [item.value for item in subjects_field]
            else:
                subjects_value = [subjects_field.value]
        elif "subjects" in form:
            subjects_value = parse_subjects_value(form.getvalue("subjects", ""))
        subject_other = form.getvalue("subjectOther", "").strip()
        if subject_other:
            subjects_value.append(subject_other)
        subjects_value = parse_subjects_value(subjects_value)
        metadata = {
            "filename": filename,
            "country": form.getvalue("country", "").strip(),
            "subjects": subjects_value,
            "gradeFrom": form.getvalue("gradeFrom", "").strip(),
            "gradeTo": form.getvalue("gradeTo", "").strip(),
            "ageFrom": form.getvalue("ageFrom", "").strip(),
            "ageTo": form.getvalue("ageTo", "").strip(),
            "specifics": specifics,
            "documentText": extracted_text[:30000],
            "sourceDocument": {
                "filename": filename,
                "storedFilename": stored_filename,
                "path": stored_path,
            },
        }

        log_message(
            "[analyze] Meta country={country} subjects={subjects} grade={g1}-{g2} age={a1}-{a2}".format(
                country=metadata["country"],
                subjects=", ".join(metadata["subjects"]),
                g1=metadata["gradeFrom"],
                g2=metadata["gradeTo"],
                a1=metadata["ageFrom"],
                a2=metadata["ageTo"],
            )
        )
        if not extracted_text.strip():
            log_message("[analyze] Warning: extracted text is empty")

        analysis_request_metadata = dict(metadata)
        analysis_request_metadata.pop("sourceDocument", None)
        analysis_request_metadata["ui_target_context"] = build_ui_target_context(metadata)
        analysis = call_mistral_analysis(analysis_request_metadata)
        if isinstance(analysis, dict):
            analysis["sourceDocument"] = metadata["sourceDocument"]
        if "error" in analysis:
            log_message(f"[analyze] Failed: {analysis.get('error')}")
            self.send_response(HTTPStatus.BAD_GATEWAY)
        else:
            log_message("[analyze] Success")
            self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps(analysis, ensure_ascii=False).encode("utf-8"))

    def handle_step2_prepare(self):
        content_type = self.headers.get("Content-Type", "")
        if not content_type.startswith("multipart/form-data"):
            self.send_error(HTTPStatus.BAD_REQUEST, "Expected multipart/form-data")
            return

        form = cgi.FieldStorage(
            fp=self.rfile,
            headers=self.headers,
            environ={
                "REQUEST_METHOD": "POST",
                "CONTENT_TYPE": content_type,
            },
        )

        country = form.getvalue("country", "").strip()
        if not country:
            self.send_error(HTTPStatus.BAD_REQUEST, "Missing country")
            return

        subjects_value = []
        if "subjects[]" in form:
            subjects_field = form["subjects[]"]
            if isinstance(subjects_field, list):
                subjects_value = [item.value for item in subjects_field]
            else:
                subjects_value = [subjects_field.value]
        elif "subjects" in form:
            subjects_field = form["subjects"]
            if isinstance(subjects_field, list):
                subjects_value = [item.value for item in subjects_field]
            else:
                subjects_value = [subjects_field.value]
        else:
            parsed_subjects = parse_subjects_value(form.getvalue("subjects", ""))
            subjects_value = parsed_subjects

        folder_info = ensure_step2_folder_structure(country, subjects_value)
        session_id = uuid.uuid4().hex[:12]
        session_root = os.path.join(folder_info["amendmentsRoot"], f"step2-{session_id}")
        os.makedirs(session_root, exist_ok=True)

        source_document = None
        raw_source_document = form.getvalue("sourceDocument", "")
        if raw_source_document:
            try:
                source_document = normalize_source_document(json.loads(raw_source_document))
            except Exception:
                source_document = normalize_source_document(raw_source_document)

        saved_files = []
        if "additionalFiles" in form:
            additional_field = form["additionalFiles"]
            file_items = additional_field if isinstance(additional_field, list) else [additional_field]
            saved_files = save_uploaded_files_to_folder(file_items, session_root)

        metadata_path = os.path.join(session_root, "step2-metadata.json")
        metadata = {
            "country": folder_info["country"],
            "subjects": [item["subject"] for item in folder_info["subjects"]],
            "sessionId": session_id,
            "createdPaths": folder_info["createdPaths"],
            "savedFiles": saved_files,
            "gradeFrom": form.getvalue("gradeFrom", "").strip(),
            "gradeTo": form.getvalue("gradeTo", "").strip(),
            "ageFrom": form.getvalue("ageFrom", "").strip(),
            "ageTo": form.getvalue("ageTo", "").strip(),
            "specifics": form.getvalue("specifics", "").strip(),
            "selections": form.getvalue("selections", "").strip(),
            "filename": form.getvalue("filename", "").strip(),
            "sourceDocument": source_document,
        }
        try:
            with open(metadata_path, "w", encoding="utf-8") as metadata_file:
                json.dump(metadata, metadata_file, ensure_ascii=False, indent=2)
        except OSError:
            self.send_error(HTTPStatus.INTERNAL_SERVER_ERROR, "Unable to save step 2 metadata")
            return

        log_message(
            "[step2] Prepared country={country} subjects={subjects} files={count}".format(
                country=folder_info["country"],
                subjects=", ".join(item["subject"] for item in folder_info["subjects"]),
                count=len(saved_files),
            )
        )

        response_body = {
            "prepared": True,
            "sessionId": session_id,
            "country": folder_info["country"],
            "countrySlug": folder_info["countrySlug"],
            "subjects": folder_info["subjects"],
            "amendmentsRoot": folder_info["amendmentsRoot"],
            "curriculumRoot": folder_info["curriculumRoot"],
            "savedFiles": saved_files,
            "filename": metadata["filename"],
            "sourceDocument": source_document,
        }

        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps(response_body, ensure_ascii=False).encode("utf-8"))

    def handle_export_docx(self):
        try:
            length = int(self.headers.get("Content-Length", 0))
            body = self.rfile.read(length)
            data = json.loads(body.decode("utf-8"))
        except Exception:
            self.send_error(HTTPStatus.BAD_REQUEST, "Invalid JSON")
            return

        try:
            context = build_export_context(data)
            bundle = build_export_bundle(context)
            export_path = bundle["docxPath"]
            with open(export_path, "rb") as export_file:
                docx_bytes = export_file.read()
        except Exception:
            log_message("[export] DOCX generation failed")
            traceback.print_exc()
            self.send_error(HTTPStatus.INTERNAL_SERVER_ERROR, "Unable to generate DOCX")
            return

        export_name = os.path.basename(export_path)
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.send_header("Content-Disposition", f'attachment; filename="{export_name}"')
        self.send_header("Content-Length", str(len(docx_bytes)))
        self.end_headers()
        self.wfile.write(docx_bytes)

    def handle_export_pdf(self):
        try:
            length = int(self.headers.get("Content-Length", 0))
            body = self.rfile.read(length)
            data = json.loads(body.decode("utf-8"))
        except Exception:
            self.send_error(HTTPStatus.BAD_REQUEST, "Invalid JSON")
            return

        try:
            context = build_export_context(data)
            bundle = build_export_bundle(context)
            export_path = bundle["pdfPath"]
            with open(export_path, "rb") as export_file:
                pdf_bytes = export_file.read()
        except Exception:
            log_message("[export] PDF generation failed")
            traceback.print_exc()
            self.send_error(HTTPStatus.INTERNAL_SERVER_ERROR, "Unable to generate PDF")
            return

        export_name = os.path.basename(export_path)
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/pdf")
        self.send_header("Content-Disposition", f'attachment; filename="{export_name}"')
        self.send_header("Content-Length", str(len(pdf_bytes)))
        self.end_headers()
        self.wfile.write(pdf_bytes)

    def handle_step3_refine(self):
        try:
            length = int(self.headers.get("Content-Length", 0))
            body = self.rfile.read(length)
            payload = json.loads(body.decode("utf-8"))
        except Exception:
            self.send_error(HTTPStatus.BAD_REQUEST, "Invalid JSON")
            return

        if not isinstance(payload, dict):
            self.send_error(HTTPStatus.BAD_REQUEST, "Missing payload")
            return

        focus_id = ""
        current_node = payload.get("currentNode")
        if isinstance(current_node, dict):
            focus_id = str(current_node.get("id", "")).strip()

        session_id = str(payload.get("sessionId", "")).strip()
        log_message(f"[mistral-step3] Refinement request session={session_id or 'unknown'} focus={focus_id or 'unknown'}")

        result = call_mistral_refinement(payload)
        session_id = str(payload.get("sessionId", "")).strip()
        if session_id:
            saved_conversation = list(payload.get("conversation", [])) if isinstance(payload.get("conversation", []), list) else []
            assistant_message = str(result.get("assistant_message", "")).strip() if isinstance(result, dict) else ""
            if assistant_message:
                saved_conversation.append({
                    "role": "tutor",
                    "text": assistant_message,
                    "time": payload.get("timestamp", ""),
                    "extraClass": ""
                })
            state_snapshot = {
                "sessionId": session_id,
                "draft": result.get("updated_draft") if isinstance(result, dict) and isinstance(result.get("updated_draft"), dict) else payload.get("draft"),
                "conversation": saved_conversation,
                "progress": payload.get("progress") if isinstance(payload.get("progress"), dict) else {},
                "currentNodeId": payload.get("currentNode", {}).get("id") if isinstance(payload.get("currentNode"), dict) else "",
                "sourceDocument": payload.get("sourceDocument") if isinstance(payload.get("sourceDocument"), dict) else None,
                "payload": payload,
                "result": result,
            }
            save_step3_state(session_id, state_snapshot)
            save_step3_conversation_log(session_id, saved_conversation, payload, result)
        if "error" in result:
            log_message(f"[mistral-step3] Failed: {result.get('error')}")
            self.send_response(HTTPStatus.BAD_GATEWAY)
        else:
            log_message("[mistral-step3] Success")
            self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps(result, ensure_ascii=False).encode("utf-8"))

    def handle_step3_state(self):
        session_id = ""
        try:
            query = urllib.parse.urlparse(self.path).query
            params = urllib.parse.parse_qs(query)
            session_id = str(params.get("sessionId", [""])[0]).strip()
        except Exception:
            session_id = ""

        if not session_id:
            self.send_error(HTTPStatus.BAD_REQUEST, "Missing sessionId")
            return

        state = load_step3_state(session_id)
        if not state:
            self.send_error(HTTPStatus.NOT_FOUND, "Step 3 state not found")
            return

        body = json.dumps(state, ensure_ascii=False).encode("utf-8")
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_send_email(self):
        try:
            length = int(self.headers.get("Content-Length", 0))
            body = self.rfile.read(length)
            payload = json.loads(body.decode("utf-8"))
        except Exception:
            self.send_error(HTTPStatus.BAD_REQUEST, "Invalid JSON")
            return

        if not isinstance(payload, dict):
            self.send_error(HTTPStatus.BAD_REQUEST, "Missing payload")
            return

        result = send_summary_email(payload)
        if "error" in result:
            log_message(f"[mail] Failed: {result.get('error')}")
            self.send_response(HTTPStatus.BAD_GATEWAY)
        else:
            log_message(f"[mail] Sent to {', '.join(result.get('recipient', []))}")
            self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps(result, ensure_ascii=False).encode("utf-8"))

    def handle_mistral_test(self):
        log_message("[mistral-test] Starting hello check")
        result = call_mistral_hello()
        if "error" in result:
            log_message(f"[mistral-test] Failed: {result.get('error')}")
            self.send_response(HTTPStatus.BAD_GATEWAY)
        else:
            log_message("[mistral-test] Success")
            self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps(result, ensure_ascii=False).encode("utf-8"))

    def handle_client_log(self):
        try:
            length = int(self.headers.get("Content-Length", 0))
            body = self.rfile.read(length)
            data = json.loads(body.decode("utf-8"))
        except Exception:
            self.send_error(HTTPStatus.BAD_REQUEST, "Invalid JSON")
            return

        message = ""
        if isinstance(data, dict):
            message = str(data.get("message", "")).strip()

        if not message:
            self.send_error(HTTPStatus.BAD_REQUEST, "Missing message")
            return

        log_message(f"[client] {message}")
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps({"saved": True}).encode("utf-8"))

    def serve_config(self):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                content = f.read()
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content.encode("utf-8"))
        except OSError:
            self.send_error(HTTPStatus.NOT_FOUND, "Config file not found")

    def handle_config_update(self):
        try:
            length = int(self.headers.get("Content-Length", 0))
            body = self.rfile.read(length)
            data = json.loads(body.decode("utf-8"))
        except Exception:
            self.send_error(HTTPStatus.BAD_REQUEST, "Invalid JSON")
            return

        if not isinstance(data, dict) or ("newSubject" not in data and "newCountry" not in data):
            self.send_error(HTTPStatus.BAD_REQUEST, "Missing newSubject or newCountry")
            return

        new_subject = None
        new_country = None
        if "newSubject" in data:
            new_subject = str(data["newSubject"]).strip()
        if "newCountry" in data:
            new_country = str(data["newCountry"]).strip()

        if new_subject is not None and not new_subject:
            self.send_error(HTTPStatus.BAD_REQUEST, "Invalid subject")
            return
        if new_country is not None and not new_country:
            self.send_error(HTTPStatus.BAD_REQUEST, "Invalid country")
            return

        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                config = json.load(f)
        except Exception:
            config = {"countries": [], "subjects": []}

        saved_item = None
        if new_subject is not None:
            if new_subject not in config.get("subjects", []):
                config.setdefault("subjects", []).append(new_subject)
                saved_item = new_subject
        if new_country is not None:
            if new_country not in config.get("countries", []):
                config.setdefault("countries", []).append(new_country)
                saved_item = new_country

        if saved_item is not None:
            with open(CONFIG_FILE, "w", encoding="utf-8") as f:
                json.dump(config, f, indent=2)

        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps({"saved": saved_item}).encode("utf-8"))

    def serve_path(self, rel_path, root):
        rel_path = rel_path.lstrip("/")
        safe_path = posixpath.normpath(urllib.parse.unquote(rel_path))
        if safe_path.startswith(".."):
            self.send_error(HTTPStatus.FORBIDDEN, "Forbidden")
            return

        file_path = os.path.join(root, safe_path)
        if os.path.isdir(file_path):
            self.send_error(HTTPStatus.FORBIDDEN, "Directory listing is not allowed")
            return
        if not os.path.exists(file_path):
            self.send_error(HTTPStatus.NOT_FOUND, "File not found")
            return
        self.serve_file(file_path, self.guess_type(file_path))

    def serve_file(self, file_path, content_type):
        try:
            with open(file_path, "rb") as f:
                content = f.read()
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)
        except OSError:
            self.send_error(HTTPStatus.NOT_FOUND, "File not found")

    def guess_type(self, path):
        if path.endswith(".html"):
            return "text/html"
        if path.endswith(".css"):
            return "text/css"
        if path.endswith(".js"):
            return "application/javascript"
        if path.endswith(".json"):
            return "application/json"
        if path.endswith(".png"):
            return "image/png"
        if path.endswith(".jpg") or path.endswith(".jpeg"):
            return "image/jpeg"
        if path.endswith(".svg"):
            return "image/svg+xml"
        if path.endswith(".pdf"):
            return "application/pdf"
        return "application/octet-stream"

def run_server():
    os.chdir(APP_ROOT)
    server_address = ("", PORT)
    httpd = http.server.HTTPServer(server_address, TutorHandler)
    url = f"http://localhost:{PORT}"
    webbrowser.open(url)
    log_message(f"Starting server at {url}")
    httpd.serve_forever()

if __name__ == "__main__":
    run_server()
